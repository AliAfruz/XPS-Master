"""Workbook import and publication rendering for XPS-master Manuscript Plotter.

This module is adapted from the user's XPS Plot Studio 1.5 core.  Its importer
keeps the original Thermo Avantage-style mapping while also understanding the
auditable multi-sheet workbooks exported by XPS-master.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any
import json
import math
import os
import re

import numpy as np
from matplotlib import colors as mcolors
from matplotlib.figure import Figure
from matplotlib.path import Path as MplPath
from matplotlib.patches import PathPatch
from openpyxl import load_workbook


PALETTE = [
    "#ff4b4b", "#6c63ff", "#2ecc71", "#f39c12", "#00bcd4",
    "#ffeb3b", "#9c27b0", "#ff7f50", "#3f51b5", "#8bc34a",
]
MIN_EXPORT_DPI = 600


def safe_filename_part(value: str) -> str:
    """Return a portable filename fragment while keeping it recognizable."""
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", value).strip(" ._")
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned or "plot"


def planned_export_paths(base_path: str | Path, sheet_names: list[str]) -> dict[str, Path]:
    """Build unique per-sheet output paths from a user-selected base path."""
    base = Path(base_path)
    suffix = base.suffix or ".png"
    root = base.with_suffix("")
    used: set[str] = set()
    result: dict[str, Path] = {}
    for sheet_name in sheet_names:
        fragment = safe_filename_part(sheet_name)
        candidate = root.parent / f"{root.name}_{fragment}{suffix}"
        key = os.path.normcase(str(candidate))
        counter = 2
        while key in used:
            candidate = root.parent / f"{root.name}_{fragment}_{counter}{suffix}"
            key = os.path.normcase(str(candidate))
            counter += 1
        used.add(key)
        result[sheet_name] = candidate
    return result


def save_figure(figure: Figure, path: str | Path, dpi: int, transparent: bool) -> None:
    """Save at publication resolution without changing canvas geometry."""

    effective_dpi = max(MIN_EXPORT_DPI, int(dpi))
    figure.savefig(
        path,
        dpi=effective_dpi,
        transparent=transparent,
        facecolor=figure.get_facecolor(),
        edgecolor=figure.get_edgecolor(),
        bbox_inches=None,
        pad_inches=0,
    )


@dataclass
class TraceStyle:
    key: str
    name: str
    kind: str
    visible: bool = True
    color: str = "#222222"
    line_width: float = 1.0
    line_style: str = "-"
    alpha: float = 1.0
    fill: bool = False
    fill_alpha: float = 0.55
    gradient_ratio: float = 0.55
    zorder: float = 3.0


@dataclass
class AnnotationSpec:
    text: str
    x: float
    y: float
    rotation: float = 90.0
    font_size: float = 12.0
    color: str = "#111111"
    bold: bool = False
    horizontal_alignment: str = "center"
    vertical_alignment: str = "bottom"


@dataclass
class PlotSettings:
    title: str = ""
    x_label: str = "Binding energy (eV)"
    y_label: str = "Intensity (count)"
    font_family: str = "Times New Roman"
    title_size: float = 20.4
    label_size: float = 19.0
    tick_size: float = 14.0
    axis_width: float = 1.8
    tick_width: float = 1.5
    tick_length: float = 6.0
    minor_tick_length: float = 3.5
    figure_width: float = 7.2
    figure_height: float = 6.6
    dpi: int = MIN_EXPORT_DPI
    reverse_x: bool = True
    square_axes: bool = True
    show_y_tick_labels: bool = False
    show_minor_ticks: bool = True
    show_legend: bool = False
    show_grid: bool = False
    transparent_export: bool = False
    show_residual_panel: bool = True
    subtract_background: bool = True
    normalization: str = "None"
    x_min: float | None = None
    x_max: float | None = None
    y_min: float | None = 0.0
    y_max: float | None = None
    background_color: str = "#ffffff"
    margins_left: float = 0.16
    margins_right: float = 0.97
    margins_bottom: float = 0.18
    margins_top: float = 0.94


@dataclass
class XPSData:
    source_path: str
    sheet_name: str
    energy: np.ndarray
    measured: np.ndarray
    envelope: np.ndarray | None = None
    background: np.ndarray | None = None
    residual: np.ndarray | None = None
    components: dict[str, np.ndarray] = field(default_factory=dict)
    column_summary: dict[str, str] = field(default_factory=dict)
    source_format: str = "Generic XPS Excel"
    components_include_background: bool = True


def list_xlsx_sheets(path: str | Path) -> list[str]:
    """Return only worksheets that contain a recognizable XPS data table."""

    # Normal mode is substantially faster than read-only mode for the small,
    # wide XPS exports this app targets (read-only random cell access is costly).
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        sheets: list[str] = []
        for name in wb.sheetnames:
            if name.casefold() in {"export summary", "global series"}:
                continue
            try:
                _find_header_row(wb[name])
            except ValueError:
                continue
            sheets.append(name)
        if not sheets:
            raise ValueError("The workbook contains no recognizable XPS data worksheets.")
        return sheets
    finally:
        wb.close()


def _numeric(value: Any) -> float:
    if value is None or value == "":
        return math.nan
    try:
        return float(value)
    except (TypeError, ValueError):
        return math.nan


def _find_header_row(ws) -> int:
    for row in range(1, min(ws.max_row, 160) + 1):
        labels = [ws.cell(row, col).value for col in range(1, min(ws.max_column, 256) + 1)]
        nonempty = [value for value in labels if value not in (None, "")]
        # The export decorative subtitle says "Binding energy" in one merged
        # cell. A scientific table header necessarily names another column too.
        if len(nonempty) < 2:
            continue
        for value in labels:
            if isinstance(value, str) and any(
                token in value.lower() for token in ("binding energy", "kinetic energy")
            ):
                return row
    raise ValueError("Could not find a binding- or kinetic-energy data header in this sheet.")


def load_xps_sheet(path: str | Path, sheet_name: str) -> XPSData:
    """Load Avantage or XPS-master Excel data without changing the workbook."""
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' is not present in the workbook.")
        ws = wb[sheet_name]
        header_row = _find_header_row(ws)
        possible_unit_row = header_row + 1

        header: dict[int, str] = {}
        units: dict[int, str] = {}
        for col in range(1, ws.max_column + 1):
            h = ws.cell(header_row, col).value
            u = ws.cell(possible_unit_row, col).value
            header[col] = str(h).strip() if h is not None else ""
            units[col] = str(u).strip() if u is not None else ""

        def cols_containing(*tokens: str) -> list[int]:
            result = []
            for col, label in header.items():
                low = label.lower()
                if any(token in low for token in tokens):
                    result.append(col)
            return result

        energy_cols = cols_containing("binding energy", "kinetic energy")
        if not energy_cols:
            raise ValueError("No binding- or kinetic-energy column was detected.")
        energy_col = next(
            (col for col in energy_cols if "active" in header[col].lower()),
            next((col for col in energy_cols if "original" not in header[col].lower()), energy_cols[0]),
        )

        def accepted_columns(*tokens: str) -> list[int]:
            return [
                col
                for col in cols_containing(*tokens)
                if "compare ·" not in header[col].lower() and "compare •" not in header[col].lower()
            ]

        envelope_cols = accepted_columns("envelope")
        background_cols = [
            col
            for col in accepted_columns("backgnd", "background")
            if "+ background" not in header[col].lower()
        ]
        residual_cols = [
            col
            for col in accepted_columns("residual")
            if "weighted" not in header[col].lower()
        ]
        measured_cols = accepted_columns("measured intensity", "measured spectrum")

        # XPS-master exports both a net component and the same component plus its
        # background.  Prefer the latter: it matches Avantage's fitted-peak
        # convention and lets the existing subtract-background switch work.
        total_component_cols = [
            col
            for col, label in header.items()
            if "+ background" in label.lower() and "compare ·" not in label.lower()
        ]
        net_component_cols = [
            col
            for col, label in header.items()
            if re.search(r"(?:—|–|-)\s*net\s*$", label, flags=re.IGNORECASE)
            and "net measured" not in label.lower()
            and "compare ·" not in label.lower()
        ]
        legacy_peak_cols = accepted_columns("fitted peak", "component")
        if total_component_cols:
            peak_cols = total_component_cols
            components_include_background = True
        elif net_component_cols:
            peak_cols = net_component_cols
            components_include_background = False
        else:
            peak_cols = legacy_peak_cols
            components_include_background = True

        first_energy_value = _numeric(ws.cell(header_row + 1, energy_col).value)
        data_start_row = header_row + 1 if math.isfinite(first_energy_value) else header_row + 2
        excluded = {
            *energy_cols,
            *envelope_cols,
            *background_cols,
            *residual_cols,
            *peak_cols,
            *cols_containing("weighted residual", "net measured", "compare ·", "compare •"),
        }
        numeric_candidates: list[int] = []
        probe_rows = range(data_start_row, min(ws.max_row, data_start_row + 28) + 1)
        for col in range(1, ws.max_column + 1):
            if col in excluded:
                continue
            score = sum(math.isfinite(_numeric(ws.cell(r, col).value)) for r in probe_rows)
            if score >= 5:
                numeric_candidates.append(col)
        measured_col = measured_cols[0] if measured_cols else (numeric_candidates[0] if numeric_candidates else None)
        if measured_col is None:
            raise ValueError("No measured-intensity column was detected.")

        selected = [energy_col, measured_col, *envelope_cols[:1], *background_cols[:1], *residual_cols[:1], *peak_cols]
        rows: list[list[float]] = []
        empty_run = 0
        for row in range(data_start_row, ws.max_row + 1):
            values = [_numeric(ws.cell(row, col).value) for col in selected]
            if not math.isfinite(values[0]):
                # XPS channels form one contiguous energy table. Stop at its
                # first terminator so later parameter/audit tables whose first
                # column is numeric are never appended as spectrum points.
                if rows:
                    break
                empty_run += 1
                continue
            empty_run = 0
            rows.append(values)
        if len(rows) < 3:
            raise ValueError("The selected sheet does not contain enough numeric XPS points.")

        matrix = np.asarray(rows, dtype=float)
        energy = matrix[:, 0]
        measured = matrix[:, 1]
        cursor = 2
        envelope = matrix[:, cursor] if envelope_cols else None
        cursor += int(bool(envelope_cols))
        background = matrix[:, cursor] if background_cols else None
        cursor += int(bool(background_cols))
        residual = matrix[:, cursor] if residual_cols else None
        cursor += int(bool(residual_cols))

        components: dict[str, np.ndarray] = {}
        for index, col in enumerate(peak_cols):
            raw_name = header[col] or f"Component {index + 1}"
            name = re.sub(r"\s*\+\s*background\s*$", "", raw_name, flags=re.IGNORECASE)
            name = re.sub(r"\s*(?:—|–|-)\s*net\s*$", "", name, flags=re.IGNORECASE)
            name = re.sub(r"fitted peak", "", name, flags=re.IGNORECASE).strip()
            if name.lower().startswith(sheet_name.lower()):
                name = name[len(sheet_name):].strip()
            if not name:
                name = f"Component {index + 1}"
            if name in components:
                name = f"{name} {index + 1}"
            components[name] = matrix[:, cursor + index]

        source_format = (
            "XPS-master"
            if any("active " in label.lower() and " energy" in label.lower() for label in header.values())
            else "Thermo Avantage / compatible"
        )
        summary = {
            "energy": f"{_excel_col(energy_col)}: {header[energy_col]}",
            "measured": f"{_excel_col(measured_col)}: {header[measured_col] or units[measured_col] or 'Measured intensity'}",
            "envelope": f"{_excel_col(envelope_cols[0])}: {header[envelope_cols[0]]}" if envelope_cols else "Not found",
            "background": f"{_excel_col(background_cols[0])}: {header[background_cols[0]]}" if background_cols else "Not found",
            "residual": f"{_excel_col(residual_cols[0])}: {header[residual_cols[0]]}" if residual_cols else "Not found",
            "components": ", ".join(_excel_col(c) for c in peak_cols) or "None",
            "points": str(len(energy)),
            "format": source_format,
        }
        return XPSData(
            source_path=str(Path(path)),
            sheet_name=sheet_name,
            energy=energy,
            measured=measured,
            envelope=envelope,
            background=background,
            residual=residual,
            components=components,
            column_summary=summary,
            source_format=source_format,
            components_include_background=components_include_background,
        )
    finally:
        wb.close()


def _excel_col(number: int) -> str:
    out = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        out = chr(65 + remainder) + out
    return out


def default_trace_styles(data: XPSData) -> list[TraceStyle]:
    styles = [TraceStyle("measured", "Measured spectrum", "Measured", True, "#4a4a4a", 0.55, "-", 0.82, False, 0.0, 0.0, 8)]
    if data.envelope is not None:
        styles.append(TraceStyle("envelope", "Fitted envelope", "Envelope", True, "#d62728", 1.15, "-", 0.92, False, 0.0, 0.0, 9))
    if data.background is not None:
        styles.append(TraceStyle("background", "Background", "Background", not True, "#7f6000", 1.0, "--", 0.9, False, 0.0, 0.0, 4))
    if data.residual is not None:
        styles.append(TraceStyle("residual", "Residual", "Residual", True, "#555555", 0.8, "-", 0.9, False, 0.0, 0.0, 2))
    for index, name in enumerate(data.components):
        styles.append(TraceStyle(f"component:{name}", name, "Component", True, PALETTE[index % len(PALETTE)], 0.8, "-", 0.95, True, 0.58, 0.58, 5 + index * 0.01))
    return styles


def _finite_xy(x: np.ndarray, y: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    mask = np.isfinite(x) & np.isfinite(y)
    return x[mask], y[mask]


def _gradient_fill(ax, x: np.ndarray, base: np.ndarray, y: np.ndarray, color: str, alpha: float, ratio: float, zorder: float):
    if np.isscalar(base):
        base = np.full_like(np.asarray(y), float(base))
    else:
        base = np.asarray(base)
    x = np.asarray(x)
    y = np.asarray(y)
    mask = np.isfinite(x) & np.isfinite(y) & np.isfinite(base)
    x, y, base = x[mask], y[mask], base[mask]
    if len(x) < 2:
        return
    ratio = float(np.clip(ratio, 0.0, 1.0))
    if ratio <= 0.001:
        ax.fill_between(x, base, y, color=color, alpha=alpha, linewidth=0, zorder=zorder)
        return
    collection = ax.fill_between(x, base, y, facecolor="none", edgecolor="none", zorder=zorder)
    vertices = np.concatenate([np.column_stack([x, y]), np.column_stack([x[::-1], base[::-1]])])
    vertices = np.vstack([vertices, vertices[0]])
    codes = np.full(len(vertices), MplPath.LINETO, dtype=np.uint8)
    codes[0] = MplPath.MOVETO
    codes[-1] = MplPath.CLOSEPOLY
    clip = PathPatch(MplPath(vertices, codes), transform=ax.transData)
    rgba = np.array(mcolors.to_rgba(color))
    gradient = np.empty((256, 1, 4), dtype=float)
    gradient[:, :, :3] = rgba[:3]
    gradient[:, :, 3] = np.linspace(alpha * (1.0 - ratio), alpha, 256)[:, None]
    ymin = float(np.nanmin(np.minimum(base, y)))
    ymax = float(np.nanmax(np.maximum(base, y)))
    if math.isclose(ymin, ymax):
        return
    image = ax.imshow(gradient, aspect="auto", origin="lower", extent=[float(np.nanmin(x)), float(np.nanmax(x)), ymin, ymax], interpolation="bicubic", zorder=zorder)
    image.set_clip_path(clip)
    collection.remove()


def transformed_series(data: XPSData, settings: PlotSettings) -> dict[str, np.ndarray]:
    bg = data.background if data.background is not None else np.zeros_like(data.measured)
    result: dict[str, np.ndarray] = {"measured": data.measured.copy()}
    if data.envelope is not None:
        result["envelope"] = data.envelope.copy()
    if data.background is not None:
        result["background"] = data.background.copy()
    if data.residual is not None:
        result["residual"] = data.residual.copy()
    for name, values in data.components.items():
        result[f"component:{name}"] = values.copy()

    if settings.subtract_background:
        for key in ("measured", "envelope"):
            if key in result:
                result[key] = result[key] - bg
        if data.components_include_background:
            for name in data.components:
                key = f"component:{name}"
                result[key] = result[key] - bg
        result["background"] = np.zeros_like(bg)

    reference = result.get("envelope", result["measured"])
    finite = reference[np.isfinite(reference)]
    scale = 1.0
    if finite.size:
        maximum = float(np.max(np.abs(finite)))
        if maximum > 0:
            if settings.normalization == "Maximum = 1":
                scale = maximum
            elif settings.normalization == "Maximum = 100":
                scale = maximum / 100.0
    if scale != 1.0:
        result = {key: value / scale for key, value in result.items()}
    return result


def make_figure(data: XPSData, settings: PlotSettings, traces: list[TraceStyle], annotations: list[AnnotationSpec]) -> tuple[Figure, Any, dict[str, Any]]:
    fig = Figure(figsize=(settings.figure_width, settings.figure_height), dpi=100, facecolor=settings.background_color)
    residual_trace = next((trace for trace in traces if trace.kind == "Residual"), None)
    show_residual = bool(
        settings.show_residual_panel
        and data.residual is not None
        and residual_trace is not None
        and residual_trace.visible
    )
    residual_ax = None
    if show_residual:
        # Explicit shared panel geometry is critical here. ``set_box_aspect``
        # adjusts each subplot independently and can make the short residual
        # panel physically wider than the square main panel even with sharex.
        # These two axes therefore receive the exact same left/right bounds;
        # sharex then keeps their limits, ticks, toolbar pan, and zoom linked.
        available_width = settings.margins_right - settings.margins_left
        available_height = settings.margins_top - settings.margins_bottom
        if available_width <= 0.0 or available_height <= 0.0:
            raise ValueError("Figure margins leave no usable manuscript plotting area.")
        residual_ratio = 0.24
        panel_gap = min(0.035, available_height * 0.08)
        main_height = (available_height - panel_gap) / (1.0 + residual_ratio)
        residual_height = residual_ratio * main_height
        axes_bottom = settings.margins_bottom
        axes_width = available_width
        axes_left = settings.margins_left
        if settings.square_axes:
            square_width = main_height * settings.figure_height / settings.figure_width
            if square_width <= available_width:
                axes_width = square_width
                axes_left += 0.5 * (available_width - axes_width)
            else:
                main_height = available_width * settings.figure_width / settings.figure_height
                residual_height = residual_ratio * main_height
                used_height = main_height + panel_gap + residual_height
                axes_bottom += 0.5 * max(0.0, available_height - used_height)
        residual_position = [axes_left, axes_bottom, axes_width, residual_height]
        main_position = [
            axes_left,
            axes_bottom + residual_height + panel_gap,
            axes_width,
            main_height,
        ]
        ax = fig.add_axes(main_position)
        residual_ax = fig.add_axes(residual_position, sharex=ax)
    else:
        ax = fig.add_subplot(111)
    ax.set_facecolor(settings.background_color)
    if residual_ax is not None:
        residual_ax.set_facecolor(settings.background_color)
    series = transformed_series(data, settings)
    artists: dict[str, Any] = {}
    baseline = series.get("background", np.zeros_like(data.energy))

    for trace in sorted(traces, key=lambda item: item.zorder):
        if trace.kind == "Residual":
            continue
        if not trace.visible or trace.key not in series:
            continue
        y = series[trace.key]
        x, clean_y = _finite_xy(data.energy, y)
        if not len(x):
            continue
        if trace.fill and trace.kind == "Component":
            base = baseline[np.isfinite(data.energy) & np.isfinite(y)] if not settings.subtract_background else np.zeros_like(clean_y)
            _gradient_fill(ax, x, base, clean_y, trace.color, trace.fill_alpha, trace.gradient_ratio, trace.zorder)
        line, = ax.plot(x, clean_y, color=trace.color, linewidth=trace.line_width, linestyle=trace.line_style, alpha=trace.alpha, label=trace.name, zorder=trace.zorder + 0.2)
        artists[trace.key] = line

    x_label_axis = residual_ax if residual_ax is not None else ax
    x_label_axis.set_xlabel(settings.x_label, fontsize=settings.label_size, fontfamily=settings.font_family, labelpad=14)
    ax.set_ylabel(settings.y_label, fontsize=settings.label_size, fontfamily=settings.font_family, labelpad=12)
    if settings.title:
        ax.set_title(settings.title, fontsize=settings.title_size, fontfamily=settings.font_family, fontweight="bold", loc="right", pad=10)
    ax.tick_params(axis="both", which="major", direction="in", width=settings.tick_width, length=settings.tick_length, labelsize=settings.tick_size, top=False, right=False)
    ax.tick_params(axis="both", which="minor", direction="in", width=max(0.5, settings.tick_width * 0.8), length=settings.minor_tick_length, top=False, right=False)
    if residual_ax is not None and residual_trace is not None:
        x, residual_y = _finite_xy(data.energy, series["residual"])
        residual_ax.axhline(0.0, color="#888888", linewidth=0.7, linestyle="--", alpha=0.8, zorder=1)
        residual_line, = residual_ax.plot(
            x,
            residual_y,
            color=residual_trace.color,
            linewidth=residual_trace.line_width,
            linestyle=residual_trace.line_style,
            alpha=residual_trace.alpha,
            label=residual_trace.name,
            zorder=3,
        )
        artists[residual_trace.key] = residual_line
        artists["residual_axis"] = residual_ax
        residual_ax.set_ylabel(
            residual_trace.name,
            fontsize=settings.label_size,
            fontfamily=settings.font_family,
            labelpad=8,
        )
        residual_ax.tick_params(
            axis="both",
            which="major",
            direction="in",
            width=settings.tick_width,
            length=settings.tick_length,
            labelsize=settings.tick_size,
            top=False,
            right=False,
        )
        residual_ax.tick_params(
            axis="both",
            which="minor",
            direction="in",
            width=max(0.5, settings.tick_width * 0.8),
            length=settings.minor_tick_length,
            top=False,
            right=False,
        )
        finite_residual = residual_y[np.isfinite(residual_y)]
        if finite_residual.size:
            amplitude = float(np.max(np.abs(finite_residual)))
            if amplitude > 0.0:
                residual_ax.set_ylim(-1.08 * amplitude, 1.08 * amplitude)
        ax.tick_params(axis="x", which="both", labelbottom=False)
    for target_ax in (ax, residual_ax):
        if target_ax is None:
            continue
        for label in target_ax.get_xticklabels() + target_ax.get_yticklabels():
            label.set_fontfamily(settings.font_family)
            label.set_fontsize(settings.tick_size)
        for offset_text in (target_ax.xaxis.get_offset_text(), target_ax.yaxis.get_offset_text()):
            offset_text.set_fontfamily(settings.font_family)
            offset_text.set_fontsize(settings.tick_size)
    if not settings.show_y_tick_labels:
        ax.tick_params(axis="y", labelleft=False)
        if residual_ax is not None:
            residual_ax.tick_params(axis="y", labelleft=False)
    if settings.show_minor_ticks:
        ax.minorticks_on()
        if residual_ax is not None:
            residual_ax.minorticks_on()
    else:
        ax.minorticks_off()
        if residual_ax is not None:
            residual_ax.minorticks_off()
    for target_ax in (ax, residual_ax):
        if target_ax is None:
            continue
        for spine in target_ax.spines.values():
            spine.set_linewidth(settings.axis_width)
            spine.set_color("#111111")
    if settings.show_grid:
        ax.grid(True, color="#d9d9d9", linewidth=0.6, alpha=0.55, zorder=0)
        if residual_ax is not None:
            residual_ax.grid(True, color="#d9d9d9", linewidth=0.6, alpha=0.45, zorder=0)
    if settings.square_axes:
        if residual_ax is None:
            ax.set_box_aspect(1)

    finite_x = data.energy[np.isfinite(data.energy)]
    if settings.x_min is not None or settings.x_max is not None:
        low = settings.x_min if settings.x_min is not None else float(np.min(finite_x))
        high = settings.x_max if settings.x_max is not None else float(np.max(finite_x))
    else:
        low, high = float(np.min(finite_x)), float(np.max(finite_x))
    if settings.reverse_x:
        ax.set_xlim(max(low, high), min(low, high))
    else:
        ax.set_xlim(min(low, high), max(low, high))
    if settings.y_min is not None or settings.y_max is not None:
        bottom, top = ax.get_ylim()
        ax.set_ylim(settings.y_min if settings.y_min is not None else bottom, settings.y_max if settings.y_max is not None else top)

    annotation_artists = []
    for index, note in enumerate(annotations):
        artist = ax.text(note.x, note.y, note.text, rotation=note.rotation, fontsize=note.font_size, color=note.color, fontfamily=settings.font_family, fontweight="bold" if note.bold else "normal", ha=note.horizontal_alignment, va=note.vertical_alignment, zorder=20, picker=True)
        artist._xps_annotation_index = index
        annotation_artists.append(artist)
    artists["annotations"] = annotation_artists

    if settings.show_legend:
        legend = ax.legend(frameon=False, fontsize=max(8, settings.tick_size - 2), loc="best")
        for text in legend.get_texts():
            text.set_fontfamily(settings.font_family)
    if residual_ax is None:
        fig.subplots_adjust(left=settings.margins_left, right=settings.margins_right, bottom=settings.margins_bottom, top=settings.margins_top)
    return fig, ax, artists


def auto_annotations(data: XPSData, settings: PlotSettings, traces: list[TraceStyle]) -> list[AnnotationSpec]:
    series = transformed_series(data, settings)
    notes: list[AnnotationSpec] = []
    for trace in traces:
        if trace.kind != "Component" or trace.key not in series:
            continue
        y = series[trace.key]
        finite = np.isfinite(data.energy) & np.isfinite(y)
        if not np.any(finite):
            continue
        indices = np.flatnonzero(finite)
        idx = indices[int(np.argmax(y[finite]))]
        notes.append(AnnotationSpec(trace.name, float(data.energy[idx]), float(y[idx]) * 1.035, 90, 11, "#111111", False))
    return notes


def project_to_json(source_path: str, sheet_name: str, settings: PlotSettings, traces: list[TraceStyle], annotations: list[AnnotationSpec]) -> str:
    payload = {
        "format": "XPS Plot Studio Project",
        "version": 1,
        "source_path": source_path,
        "sheet_name": sheet_name,
        "settings": asdict(settings),
        "traces": [asdict(item) for item in traces],
        "annotations": [asdict(item) for item in annotations],
    }
    return json.dumps(payload, indent=2, ensure_ascii=False)


def project_from_json(text: str) -> dict[str, Any]:
    payload = json.loads(text)
    if payload.get("format") != "XPS Plot Studio Project":
        raise ValueError("This is not an XPS Plot Studio project file.")
    return payload


def template_to_json(settings: PlotSettings, traces: list[TraceStyle]) -> str:
    """Serialize reusable plot styling without workbook-specific data."""
    payload = {
        "format": "XPS Plot Studio Template",
        "version": 1,
        "settings": asdict(settings),
        "traces": [asdict(item) for item in traces],
    }
    return json.dumps(payload, indent=2, ensure_ascii=False)


def template_from_json(text: str) -> dict[str, Any]:
    payload = json.loads(text)
    if payload.get("format") != "XPS Plot Studio Template":
        raise ValueError("This is not an XPS Plot Studio template file.")
    if not isinstance(payload.get("settings"), dict):
        raise ValueError("The template does not contain valid plot settings.")
    if not isinstance(payload.get("traces"), list):
        raise ValueError("The template does not contain valid trace styles.")
    return payload


def apply_trace_template(current: list[TraceStyle], saved: list[TraceStyle]) -> list[TraceStyle]:
    """Apply saved styles to the traces available in the current data set.

    Exact trace keys are preferred. Component styles also fall back to their
    component order so that a template remains useful for another XPS scan
    whose fitted components have different names.
    """
    by_key = {item.key: item for item in saved}
    saved_by_kind = {item.kind: item for item in saved if item.kind != "Component"}
    saved_components = [item for item in saved if item.kind == "Component"]
    component_index = 0
    result: list[TraceStyle] = []

    for item in current:
        style = by_key.get(item.key)
        exact_match = style is not None
        if item.kind == "Component":
            current_component_index = component_index
            component_index += 1
            if style is None and current_component_index < len(saved_components):
                style = saved_components[current_component_index]
        elif style is None:
            style = saved_by_kind.get(item.kind)

        if style is None:
            result.append(item)
            continue

        values = asdict(style)
        values["key"] = item.key
        values["kind"] = item.kind
        # Component names are chemical assignments and should not be copied to
        # a differently named component merely because it has the same order.
        if item.kind == "Component" and not exact_match:
            values["name"] = item.name
        result.append(TraceStyle(**values))
    return result

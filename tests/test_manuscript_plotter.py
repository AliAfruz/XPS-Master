from pathlib import Path

import numpy as np
import pytest
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from openpyxl import Workbook
from PIL import Image

from xps_afruz.manuscript_plotter.core import (
    PlotSettings,
    default_trace_styles,
    list_xlsx_sheets,
    load_xps_sheet,
    make_figure,
    save_figure,
    transformed_series,
)


def _write_avantage_workbook(path: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    workbook = Workbook()
    notes = workbook.active
    notes.title = "Notes"
    notes["A1"] = "Instrument notes without a spectrum table"
    sheet = workbook.create_sheet("O1s Scan")
    headers = (
        "Kinetic Energy",
        "Counts",
        "Envelope",
        "Background",
        "Residual",
        "Fitted Peak O lattice",
    )
    for column, value in enumerate(headers, start=1):
        sheet.cell(3, column, value)
        sheet.cell(4, column, "eV" if column == 1 else "counts")

    energy = np.array([960.0, 959.0, 958.0, 957.0, 956.0])
    background = np.linspace(10.0, 11.0, energy.size)
    net_component = np.array([0.2, 1.0, 3.0, 1.2, 0.1])
    envelope = background + net_component
    measured = envelope + np.array([0.1, -0.1, 0.2, -0.2, 0.05])
    residual = envelope - measured
    for row, values in enumerate(
        zip(energy, measured, envelope, background, residual, background + net_component),
        start=5,
    ):
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column, float(value))
    workbook.save(path)
    return energy, measured, net_component


def test_avantage_import_remains_supported_and_uses_residual_panel(tmp_path: Path) -> None:
    path = tmp_path / "avantage.xlsx"
    energy, measured, net_component = _write_avantage_workbook(path)

    assert list_xlsx_sheets(path) == ["O1s Scan"]
    data = load_xps_sheet(path, "O1s Scan")

    assert data.source_format == "Thermo Avantage / compatible"
    assert np.array_equal(data.energy, energy)
    assert np.allclose(data.measured, measured)
    assert list(data.components) == ["O lattice"]
    transformed = transformed_series(data, PlotSettings())
    assert np.allclose(transformed["component:O lattice"], net_component)

    settings = PlotSettings(title="O 1s")
    figure, axis, artists = make_figure(
        data,
        settings,
        default_trace_styles(data),
        [],
    )
    assert len(figure.axes) == 2
    assert axis is figure.axes[0]
    assert artists["residual_axis"] is figure.axes[1]
    assert np.allclose(artists["residual"].get_ydata(), data.residual)

    main_axis, residual_axis = figure.axes
    title_artist = next(child for child in main_axis.get_children() if getattr(child, "get_text", lambda: "")() == "O 1s")
    assert title_artist.get_fontsize() == pytest.approx(20.4)
    assert title_artist.get_fontfamily() == ["Times New Roman"]
    assert main_axis.yaxis.label.get_fontsize() == pytest.approx(19.0)
    assert residual_axis.xaxis.label.get_fontsize() == pytest.approx(19.0)
    assert residual_axis.yaxis.label.get_fontsize() == pytest.approx(19.0)
    assert residual_axis.yaxis.label.get_fontfamily() == main_axis.yaxis.label.get_fontfamily()
    for target_axis in (main_axis, residual_axis):
        for tick_label in target_axis.get_xticklabels() + target_axis.get_yticklabels():
            assert tick_label.get_fontsize() == pytest.approx(14.0)
            assert tick_label.get_fontfamily() == ["Times New Roman"]

    canvas = FigureCanvasAgg(figure)
    canvas.draw()
    main_position = main_axis.get_position()
    residual_position = residual_axis.get_position()
    assert main_position.x0 == pytest.approx(residual_position.x0, abs=1e-12)
    assert main_position.x1 == pytest.approx(residual_position.x1, abs=1e-12)
    assert main_axis.get_shared_x_axes().joined(main_axis, residual_axis)
    for energy_value in data.energy:
        main_pixel = main_axis.transData.transform((energy_value, 0.0))[0]
        residual_pixel = residual_axis.transData.transform((energy_value, 0.0))[0]
        assert main_pixel == pytest.approx(residual_pixel, abs=1e-9)

    main_axis.set_xlim(959.5, 956.5)
    canvas.draw()
    assert residual_axis.get_xlim() == pytest.approx((959.5, 956.5))
    residual_axis.set_xlim(959.0, 957.0)
    canvas.draw()
    assert main_axis.get_xlim() == pytest.approx((959.0, 957.0))


def test_residual_panel_can_be_disabled(tmp_path: Path) -> None:
    path = tmp_path / "avantage.xlsx"
    _write_avantage_workbook(path)
    data = load_xps_sheet(path, "O1s Scan")

    settings = PlotSettings(show_residual_panel=False)
    figure, _axis, artists = make_figure(data, settings, default_trace_styles(data), [])

    assert len(figure.axes) == 1
    assert "residual_axis" not in artists


def test_raster_export_enforces_at_least_600_dpi(tmp_path: Path) -> None:
    assert PlotSettings().dpi == 600
    low_request = tmp_path / "minimum-dpi.png"
    high_request = tmp_path / "higher-dpi.png"
    figure = Figure(figsize=(0.5, 0.5))

    save_figure(figure, low_request, dpi=300, transparent=False)
    save_figure(figure, high_request, dpi=900, transparent=False)

    with Image.open(low_request) as image:
        assert image.size == (300, 300)
    with Image.open(high_request) as image:
        assert image.size == (450, 450)

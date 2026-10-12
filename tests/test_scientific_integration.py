"""Scientific regression checks spanning the retained fitting workflows."""
from dataclasses import replace
import numpy as np
import pytest
from xps_afruz import BackgroundMethod, BackgroundSettings, FitModel, PeakComponent, Spectrum, fit_spectrum
from xps_afruz.fitting import smart_deconvolution as smart_module
from xps_afruz.fitting.global_series import (fit_global_series, GlobalSeriesSettings,
    _GlobalVariable, _parameter_errors_for_spectrum)
from xps_afruz.fitting.model import ParameterTie
from xps_afruz.fitting.model_comparison import ModelCandidate, ModelComparisonSettings, compare_models
from xps_afruz.fitting.peak_shapes import evaluate_peak
from xps_afruz.fitting.statistics import information_criteria
from xps_afruz.io.excel_export import export_workspace_to_excel
from xps_afruz.io.project_io import WorkspaceItemData


def test_global_center_errors_include_shift_cross_covariance_and_tie_gradient():
    a = PeakComponent('A', 0, 1, 10, profile='gaussian', vary_fwhm=False, vary_area=False)
    b = replace(a, name='B', center=10)
    model = FitModel([a, b], np.zeros(20), [ParameterTie('B.center', 'A.center', scale=2, offset=10)])
    variables = (
        _GlobalVariable('shared::A.center', 'A.center', None, 0, -2, 2),
        _GlobalVariable('s1::energy_shift', None, 1, 0, -1, 1, True),
    )
    covariance = np.array([[4.0, -3.0], [-3.0, 9.0]])
    errors = _parameter_errors_for_spectrum(model, variables, covariance, 1, GlobalSeriesSettings())
    assert errors['A.center'] == pytest.approx(np.sqrt(4+9-6))
    assert errors['B.center'] == pytest.approx(np.sqrt(4*4+9-12))
    assert errors['A.area'] == 0.0


def test_global_poisson_fits_recover_area_mles_and_use_deviance_evidence():
    x = np.linspace(-3, 3, 120)
    unit = evaluate_peak(x, profile='gaussian', center=0, fwhm=1.3, area=1)
    rng = np.random.default_rng(761)
    counts = [rng.poisson(8*unit), rng.poisson(20*unit)]
    spectra = [Spectrum(x, y, label=f'Scan {index}') for index, y in enumerate(counts)]
    component = PeakComponent('A', 0, 1.3, 10, profile='gaussian', vary_center=False, vary_fwhm=False, area_bounds=(0, 100))
    settings = GlobalSeriesSettings(weighting='poisson', allow_energy_shifts=False)
    result = fit_global_series(spectra, [np.zeros_like(x)]*2, [component], settings=settings)
    assert result.success
    assert result.global_parameter_count == result.independent_parameter_count == 2
    criteria = information_criteria(result.global_objective, result.total_samples, 2, 'poisson')
    assert result.global_aicc == pytest.approx(criteria.aicc)
    assert result.global_bic == pytest.approx(criteria.bic)
    for item, data in zip(result.spectrum_fits, counts):
        area = data.sum()/unit.sum()
        assert item.global_fit.parameter_values['A.area'] == pytest.approx(area, rel=2e-5)
        assert item.global_fit.standard_errors['A.area'] == pytest.approx(np.sqrt(area/unit.sum()), rel=2e-4)
        assert item.global_fit.objective_name == 'poisson_deviance'
        assert item.independent_fit.objective_name == 'poisson_deviance'


def test_no_ineligible_model_receives_normalized_support():
    x = np.linspace(-3, 3, 100)
    unit = evaluate_peak(x, profile='gaussian', center=0, fwhm=1.3, area=1)
    spectrum = Spectrum(x, 20*unit)
    a = PeakComponent('A', 0, 1.3, 8, profile='gaussian', vary_center=False, vary_fwhm=False)
    b = replace(a, name='B', area=12)
    baseline = BackgroundSettings('custom', {'values': np.zeros_like(x).tolist()})
    candidates = [ModelCandidate(str(i), f'Degenerate {i}', (a, replace(b, area=5+i)), (), baseline, 'Unidentifiable pair')
                  for i in range(2)]
    result = compare_models(spectrum, candidates, ModelComparisonSettings(max_nfev=200))
    assert all(not candidate.eligible for candidate in result.candidates)
    assert all(candidate.model_support == candidate.aicc_weight == candidate.bic_weight == 0 for candidate in result.candidates)
    assert 'No candidate' in result.numerical_note


def test_smart_search_excludes_nonconverged_candidates(monkeypatch):
    x = np.linspace(-3, 3, 100)
    signal = evaluate_peak(x, profile='gaussian', center=0, fwhm=1.3, area=20)
    spectrum = Spectrum(x, signal)
    component = PeakComponent('A', 0, 1.3, 20, profile='gaussian')
    stalled = replace(fit_spectrum(spectrum, FitModel([component], np.zeros_like(x))), success=False)
    monkeypatch.setattr(smart_module, 'fit_spectrum', lambda *_args, **_kwargs: stalled)
    with pytest.raises(ValueError, match='No valid fit'):
        smart_module.smart_deconvolve(spectrum, np.zeros_like(x), smart_module.SmartDeconvolutionSettings(1, candidate_trials=4))


def test_excel_export_preserves_fit_objective_rank_and_warnings(tmp_path):
    from openpyxl import load_workbook
    x = np.linspace(-3, 3, 80)
    signal = evaluate_peak(x, profile='gaussian', center=0, fwhm=1.3, area=20)
    spectrum = Spectrum(x, signal)
    component = PeakComponent('A', 0, 1.3, 20, profile='gaussian', vary_center=False, vary_fwhm=False)
    fit = fit_spectrum(spectrum, FitModel([component], np.zeros_like(x)), weighting='poisson')
    item = WorkspaceItemData(spectrum, BackgroundMethod.CUSTOM, {'values': np.zeros_like(x).tolist()}, fit_result=fit)
    path = tmp_path/'diagnostics.xlsx'
    export_workspace_to_excel(path, (item,))
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[1]]
        cells = {row[0].value: row[1].value for row in sheet.iter_rows() if len(row) > 1 and row[0].value is not None}
        assert cells['Optimization objective'] == 'poisson_deviance'
        assert cells['Jacobian rank'] == 1
        assert cells['Independent parameters'] == 1
        assert 'Low expected counts' in cells['Fit review warnings']
    finally:
        workbook.close()


@pytest.mark.gui
def test_reopened_project_displays_retained_rank_deficiency_and_undefined_condition(tmp_path):
    from PySide6.QtWidgets import QApplication
    from xps_afruz.ui.main_window import MainWindow
    application = QApplication.instance() or QApplication([])
    x = np.linspace(-3, 3, 100)
    background = np.zeros_like(x)
    a = PeakComponent('A', 0, 1.1, 10, profile='gaussian', vary_center=False, vary_fwhm=False)
    b = replace(a, name='B')
    spectrum = Spectrum(x, evaluate_peak(x, profile='gaussian', center=0, fwhm=1.1, area=30))
    result = fit_spectrum(spectrum, FitModel([a, b], background))
    window = MainWindow()
    restored = MainWindow()
    try:
        window.add_workspace_spectra([spectrum], activate=True)
        window.add_peak_row(a)
        window.add_peak_row(b)
        window.fit_result = result
        window._store_active_workspace_fit_state()
        path = tmp_path/'rank-deficient.xpsa'
        assert window._write_project(path)
        assert restored._open_project_path(path)
        assert restored.fit_result.jacobian_rank == 1
        assert 'Jacobian rank' in restored.fit_summary.text()
        assert 'Rank-deficient' in restored.fit_summary.text()
        assert 'undefined' in restored.fit_summary.text()
    finally:
        window._set_project_dirty(False)
        restored._set_project_dirty(False)
        window.close()
        restored.close()
        application.processEvents()
